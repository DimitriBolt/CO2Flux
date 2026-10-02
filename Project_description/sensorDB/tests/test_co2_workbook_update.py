"""Offline regression checks: current layout, row order and workbook preservation."""
from datetime import datetime
from pathlib import Path
import hashlib
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts"))
import update_co2_sheet as update
import run_update_east_center as partial


class FakeCursor:
    def __init__(self):
        self.calls = []
        self.closed = False

    def execute(self, query, **params):
        self.calls.append((query, params))

    def fetchone(self):
        return datetime(2024, 1, 2, 3, 4, 5), 420.0

    def close(self):
        self.closed = True


class FakeConnection:
    def __init__(self):
        self.cur = FakeCursor()
        self.closed = False

    def cursor(self):
        return self.cur

    def close(self):
        self.closed = True


def fixture():
    wb = Workbook()
    ws = wb.active
    ws.title = "CO2"
    for col, label in update.EXPECTED_HEADERS.items():
        ws[f"{col}3"] = label
    # Deliberately outside the old blocks, and in a different order.
    for row, slope, sid, air in [(8, "LEO West", 1275, True),
                                (300, "LEO West", 994, False),
                                (220, "LEO Center", 1275, True),
                                (60, "LEO East", 408, False)]:
        prefix = {"LEO West": "W", "LEO Center": "C", "LEO East": "E"}[slope]
        values = {"B": "C_CO2,air" if air else "C_CO2,basalt", "I": slope,
                  "K": "LI-COR" if air else "GMM222", "L": sid,
                  "M": f"LEO-{prefix}_4_0_1_LI-7000" if air else f"LEO-{prefix}_4_-4_1_GMM222",
                  "N": update.SLOPE_SCHEMAS[slope] + (".datavalueslicor" if air else ".datavalues"),
                  "AD": 56 if air else 9, "AG": "umol/mol" if air else "ppm",
                  "R": "umol/mol" if air else "ppm", "P": "LOCALDATETIME", "Q": "DATAVALUE"}
        for col, value in values.items():
            ws[f"{col}{row}"] = value
    ws["B106"] = "C_H2O,air"
    ws["L106"] = 1275
    ws["AA106"] = "=1+2"
    ws["AK300"] = "=SUM(1,2)"
    ws.merge_cells("A310:C310")
    ws["A310"] = "unrelated merged cell"
    wb.create_sheet("Other")["B7"] = "=CO2!L8"
    return wb


class WorkbookUpdateTests(unittest.TestCase):
    def test_current_columns_and_no_numeric_fallback(self):
        self.assertEqual(update.resolve_sensor_id({"J": 9999, "L": 1275}), 1275)
        for bad in (None, True, 1.5, "not an ID"):
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                update.resolve_sensor_id({"J": 1275, "L": bad, "O": "dv.sensorid = 1275"})

    def test_selection_independent_of_old_row_numbers(self):
        wb = fixture()
        rows = update.selected_records(wb["CO2"], {"LEO West"})
        self.assertEqual([row for row, _ in rows], [8, 300])
        wb.close()

    def test_west_air_identity_units_and_queries_preserved(self):
        wb = fixture()
        ws = wb["CO2"]
        before = {c: ws[f"{c}8"].value for c in ("L", "M", "N", "AD", "AG", "R")}
        cur = FakeCursor()
        self.assertEqual(update.update_workbook(wb, cur, slopes={"LEO West"}), 2)
        self.assertEqual(before, {c: ws[f"{c}8"].value for c in before})
        self.assertIn("leo_west.datavalueslicor", ws["AA8"].value)
        self.assertIn("dv.sensorid = 1275", ws["AA8"].value)
        self.assertIn("dv.variableid = 56", ws["AA8"].value)
        self.assertEqual(cur.calls[0][1], {"sensor_id": 1275, "variable_id": 56})
        self.assertEqual(ws["AA106"].value, "=1+2")
        self.assertEqual(ws["AK300"].value, "=SUM(1,2)")
        self.assertEqual(wb["Other"]["B7"].value, "=CO2!L8")
        wb.close()

    def test_validation_precedes_queries_and_changes(self):
        for col, bad in [("AD", 1), ("N", "leo_center.datavalueslicor"),
                         ("M", None), ("AG", None), ("AA", "=1+2")]:
            with self.subTest(col=col):
                wb = fixture()
                wb["CO2"][f"{col}8"] = bad
                cur = FakeCursor()
                with self.assertRaises(ValueError):
                    update.update_workbook(wb, cur)
                self.assertEqual(cur.calls, [])
                self.assertIsNone(wb["CO2"]["AA60"].value)
                wb.close()

    def test_duplicate_full_key_rejected_but_cross_slope_ids_allowed(self):
        wb = fixture()
        self.assertEqual(len(update.selected_records(wb["CO2"], set(update.SLOPE_SCHEMAS))), 4)
        for col in update.EXPECTED_HEADERS:
            wb["CO2"][f"{col}9"] = wb["CO2"][f"{col}8"].value
        with self.assertRaises(ValueError):
            update.selected_records(wb["CO2"], {"LEO West"})
        wb.close()

    def test_empty_series_does_not_write_no_into_availability_end(self):
        wb = fixture()
        cur = FakeCursor()
        with patch.object(cur, "fetchone", return_value=None):
            update.update_workbook(wb, cur, slopes={"LEO West"})
        self.assertEqual(wb["CO2"]["Y8"].value, "")
        self.assertTrue(wb["CO2"]["AC8"].value.startswith("No"))
        wb.close()

    def test_partial_entrypoint_preserves_west_and_other_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "fixture.xlsx"
            wb = fixture()
            before = [[cell.value for cell in wb["CO2"][i]] for i in (8, 300)]
            wb.save(path)
            wb.close()
            conn = FakeConnection()
            with patch.object(update, "connect", return_value=conn):
                partial.main(workbook_path=path)
            wb = load_workbook(path)
            self.assertEqual(before, [[cell.value for cell in wb["CO2"][i]] for i in (8, 300)])
            self.assertEqual(wb["Other"]["B7"].value, "=CO2!L8")
            self.assertIn("A310:C310", str(wb["CO2"].merged_cells))
            self.assertTrue(conn.closed and conn.cur.closed)
            self.assertFalse(any("leo_west" in query for query, _ in conn.cur.calls))
            wb.close()

    def test_real_workbook_temporary_copy_preserves_every_unrelated_cell(self):
        original = ROOT / "Sensors_Description/variables_schema.xlsx"
        digest = hashlib.sha256(original.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "copy.xlsx"
            shutil.copyfile(original, path)
            before = load_workbook(path)
            allowed_rows = {row for row, _ in update.selected_records(before["CO2"], set(update.SLOPE_SCHEMAS))}
            self.assertTrue({100, 101, 102, 205}.issubset(allowed_rows))
            conn = FakeConnection()
            with patch.object(update, "connect", return_value=conn):
                update.main(workbook_path=path)
            after = load_workbook(path)
            self.assertEqual(before.sheetnames, after.sheetnames)
            for name in before.sheetnames:
                a, b = before[name], after[name]
                self.assertEqual((a.max_row, a.max_column), (b.max_row, b.max_column))
                self.assertEqual(str(a.merged_cells), str(b.merged_cells))
                for row in a:
                    for cell in row:
                        current = b[cell.coordinate]
                        self.assertEqual(cell._style, current._style)
                        if name == "CO2" and cell.row in allowed_rows and cell.column_letter in update.UPDATE_COLUMNS:
                            continue
                        self.assertEqual(cell.value, current.value, f"{name}!{cell.coordinate}")
            before.close()
            after.close()
        self.assertEqual(hashlib.sha256(original.read_bytes()).hexdigest(), digest)


if __name__ == "__main__":
    unittest.main()
